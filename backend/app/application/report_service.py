"""Visitor report queries and secure exports."""

from __future__ import annotations

import csv
import io
import re
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import or_
from sqlalchemy.orm import Session, joinedload

from app.application.audit_service import AuditService
from app.application.auth_service import AuthContext
from app.application.site_scope_service import get_allowed_location_ids
from app.application.timezone_service import get_location_day_bounds
from app.core.config import settings
from app.domain.enums import Permission, VisitStatus, role_has_permission
from app.domain.models import Host, Location, Visit, Visitor, VisitorType


REPORT_COLUMNS = [
    ("registration_reference", "Visit Reference"),
    ("visitor_name", "Visitor"),
    ("visitor_mobile", "Mobile"),
    ("visitor_email", "Email"),
    ("company", "Company"),
    ("visitor_type", "Visitor Type"),
    ("host_name", "Host"),
    ("location_name", "Location"),
    ("source", "Source"),
    ("scheduled_start", "Expected Arrival"),
    ("scheduled_end", "Expected Departure"),
    ("created_at", "Registered At"),
    ("registered_by_name", "Registered By"),
    ("registered_by_email", "Registered By Email"),
    ("approval_status", "Approval Status"),
    ("approval_time", "Approval Time"),
    ("security_status", "Security Status"),
    ("compliance_status", "Compliance Status"),
    ("checked_in_at", "Actual Check-In"),
    ("checked_out_at", "Actual Check-Out"),
    ("visit_duration_minutes", "Visit Duration (min)"),
    ("status", "Status"),
]

ALLOWED_SORT = {
    "registered_at": Visit.created_at,
    "scheduled_start": Visit.scheduled_start,
    "scheduled_end": Visit.scheduled_end,
    "checked_in_at": Visit.checked_in_at,
    "checked_out_at": Visit.checked_out_at,
    "visitor_name": Visitor.full_name,
    "company": Visitor.company,
    "location": Location.name,
    "status": Visit.status,
}


class ReportError(Exception):
    def __init__(self, code: str, message: str, status_code: int = 400):
        self.code = code
        self.message = message
        self.status_code = status_code
        super().__init__(message)


def _require_read(ctx: AuthContext) -> None:
    if not role_has_permission(ctx.role, Permission.REPORTS_VIEW):
        raise ReportError("forbidden", "Permission denied.", 403)


def _require_export_csv(ctx: AuthContext) -> None:
    if not (
        role_has_permission(ctx.role, Permission.REPORTS_EXPORT_CSV)
        or role_has_permission(ctx.role, Permission.REPORTS_EXPORT)
    ):
        raise ReportError("forbidden", "Permission denied.", 403)


def _require_export_xlsx(ctx: AuthContext) -> None:
    if not (
        role_has_permission(ctx.role, Permission.REPORTS_EXPORT_XLSX)
        or role_has_permission(ctx.role, Permission.REPORTS_EXPORT)
    ):
        raise ReportError("forbidden", "Permission denied.", 403)


def _resolve_locations(db: Session, ctx: AuthContext, location_id: Optional[int]) -> List[int]:
    allowed = get_allowed_location_ids(ctx, db, location_id)
    if allowed is None:
        if location_id is not None:
            return [location_id]
        return [r[0] for r in db.query(Location.id).filter(Location.is_active.is_(True)).all()]
    if not allowed:
        raise ReportError("forbidden", "No accessible locations.", 403)
    return allowed


def _validate_range(from_dt: datetime, to_dt: datetime) -> None:
    if to_dt <= from_dt:
        raise ReportError("invalid_range", "End date must be after start date.")
    if (to_dt - from_dt).days > settings.analytics_max_range_days:
        raise ReportError("range_too_large", f"Maximum range is {settings.analytics_max_range_days} days.")


def _build_row(visit: Visit) -> Dict[str, Any]:
    visitor = visit.visitor
    loc = visit.location
    vt_name = visitor.visitor_type.name if visitor and visitor.visitor_type else None
    creator = visit.created_by
    approval = None
    for a in sorted(visit.approvals or [], key=lambda x: x.created_at or datetime.min.replace(tzinfo=timezone.utc)):
        if a.decision == "APPROVED":
            approval = a
            break
    approval_time = None
    if approval and visit.created_at and approval.created_at:
        approval_time = approval.created_at.isoformat()
    duration_min = None
    if visit.checked_in_at and visit.checked_out_at:
        duration_min = round((visit.checked_out_at - visit.checked_in_at).total_seconds() / 60, 1)
    approval_status = "APPROVED" if approval else ("REJECTED" if visit.status == VisitStatus.REJECTED.value else visit.status)
    return {
        "registration_reference": visit.registration_reference,
        "visitor_name": visitor.full_name if visitor else "",
        "visitor_mobile": visitor.phone if visitor else None,
        "visitor_email": visitor.email if visitor else None,
        "company": visitor.company if visitor else None,
        "visitor_type": vt_name,
        "host_name": visit.host_name,
        "location_name": loc.name if loc else None,
        "source": visit.source,
        "scheduled_start": visit.scheduled_start.isoformat() if visit.scheduled_start else None,
        "scheduled_end": visit.scheduled_end.isoformat() if visit.scheduled_end else None,
        "created_at": visit.created_at.isoformat() if visit.created_at else None,
        "registered_by_name": creator.display_name if creator else None,
        "registered_by_email": creator.email if creator else None,
        "approval_status": approval_status,
        "approval_time": approval_time,
        "security_status": visit.security_clearance_status,
        "compliance_status": visit.compliance_status,
        "checked_in_at": visit.checked_in_at.isoformat() if visit.checked_in_at else None,
        "checked_out_at": visit.checked_out_at.isoformat() if visit.checked_out_at else None,
        "visit_duration_minutes": duration_min,
        "status": visit.status,
    }


def _apply_filters(
    query,
    registration_reference: Optional[str] = None,
    company: Optional[str] = None,
    host_id: Optional[int] = None,
    source: Optional[str] = None,
    status: Optional[str] = None,
    security_status: Optional[str] = None,
    compliance_status: Optional[str] = None,
    visitor_type_id: Optional[int] = None,
):
    if registration_reference:
        query = query.filter(Visit.registration_reference.ilike(f"%{registration_reference.strip()}%"))
    if company:
        query = query.filter(Visitor.company.ilike(f"%{company.strip()}%"))
    if host_id:
        query = query.filter(Visit.host_id == host_id)
    if source:
        query = query.filter(Visit.source == source)
    if status:
        query = query.filter(Visit.status == status.upper())
    if security_status:
        query = query.filter(Visit.security_clearance_status == security_status.upper())
    if compliance_status:
        query = query.filter(Visit.compliance_status == compliance_status.upper())
    if visitor_type_id:
        query = query.filter(Visitor.visitor_type_id == visitor_type_id)
    return query


def list_visit_reports(
    db: Session,
    ctx: AuthContext,
    from_dt: datetime,
    to_dt: datetime,
    location_id: Optional[int] = None,
    visitor_type_id: Optional[int] = None,
    company: Optional[str] = None,
    host_id: Optional[int] = None,
    source: Optional[str] = None,
    status: Optional[str] = None,
    security_status: Optional[str] = None,
    compliance_status: Optional[str] = None,
    registration_reference: Optional[str] = None,
    page: int = 1,
    page_size: int = 50,
    sort: str = "registered_at",
    sort_dir: str = "desc",
) -> Tuple[List[Dict[str, Any]], int]:
    _require_read(ctx)
    _validate_range(from_dt, to_dt)
    location_ids = _resolve_locations(db, ctx, location_id)
    page_size = min(max(page_size, 1), 100)
    page = max(page, 1)

    query = (
        db.query(Visit)
        .join(Visitor)
        .join(Location, Visit.location_id == Location.id)
        .options(
            joinedload(Visit.visitor).joinedload(Visitor.visitor_type),
            joinedload(Visit.location),
            joinedload(Visit.approvals),
            joinedload(Visit.created_by),
        )
        .filter(
            Visit.location_id.in_(location_ids),
            Visit.created_at >= from_dt,
            Visit.created_at < to_dt,
        )
    )
    query = _apply_filters(
        query, registration_reference, company, host_id, source, status,
        security_status, compliance_status, visitor_type_id,
    )
    sort_col = ALLOWED_SORT.get(sort, Visit.created_at)
    if sort_dir.lower() == "asc":
        query = query.order_by(sort_col.asc())
    else:
        query = query.order_by(sort_col.desc())

    total = query.count()
    visits = query.offset((page - 1) * page_size).limit(page_size).all()
    return [_build_row(v) for v in visits], total


def _fetch_export_rows(db: Session, ctx: AuthContext, filters: Dict[str, Any]) -> List[Dict[str, Any]]:
    from_dt = filters["from_dt"]
    to_dt = filters["to_dt"]
    location_ids = _resolve_locations(db, ctx, filters.get("location_id"))
    query = (
        db.query(Visit)
        .join(Visitor)
        .join(Location, Visit.location_id == Location.id)
        .options(
            joinedload(Visit.visitor).joinedload(Visitor.visitor_type),
            joinedload(Visit.location),
            joinedload(Visit.approvals),
            joinedload(Visit.created_by),
        )
        .filter(Visit.location_id.in_(location_ids), Visit.created_at >= from_dt, Visit.created_at < to_dt)
    )
    query = _apply_filters(
        query,
        filters.get("registration_reference"),
        filters.get("company"),
        filters.get("host_id"),
        filters.get("source"),
        filters.get("status"),
        filters.get("security_status"),
        filters.get("compliance_status"),
        filters.get("visitor_type_id"),
    )
    query = query.order_by(Visit.created_at.desc())
    count = query.count()
    if count > settings.report_export_max_rows:
        raise ReportError(
            "export_too_large",
            f"Export exceeds maximum of {settings.report_export_max_rows} rows. Narrow your filters.",
            400,
        )
    return [_build_row(v) for v in query.all()]


def sanitize_csv_cell(value: Any) -> str:
    if value is None:
        return ""
    s = str(value)
    if s and s[0] in ("=", "+", "-", "@"):
        return "'" + s
    return s


def build_csv(rows: List[Dict[str, Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([label for _, label in REPORT_COLUMNS])
    for row in rows:
        writer.writerow([sanitize_csv_cell(row.get(key)) for key, _ in REPORT_COLUMNS])
    return buf.getvalue().encode("utf-8-sig")


def build_xlsx(rows: List[Dict[str, Any]], summary: Optional[Dict[str, Any]] = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Visitor Report"
    headers = [label for _, label in REPORT_COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(key) for key, _ in REPORT_COLUMNS])
    ws.freeze_panes = "A2"
    if summary:
        ws2 = wb.create_sheet("Report Summary")
        for k, v in summary.items():
            ws2.append([k, v])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def safe_filename_part(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", s.strip())[:40] or "all"


def export_visits_csv(db: Session, ctx: AuthContext, filters: Dict[str, Any]) -> Tuple[bytes, str, int]:
    _require_export_csv(ctx)
    rows = _fetch_export_rows(db, ctx, filters)
    AuditService(db).record(
        action="REPORT_EXPORTED_CSV",
        entity_type="report",
        entity_id="visits",
        actor_id=ctx.user_id,
        actor_email=ctx.email,
        metadata={"row_count": len(rows), "filters": {k: v for k, v in filters.items() if k not in ("from_dt", "to_dt")}},
    )
    db.commit()
    loc_part = safe_filename_part(str(filters.get("location_label") or "all"))
    from_s = filters["from_dt"].strftime("%Y-%m-%d")
    to_s = filters["to_dt"].strftime("%Y-%m-%d")
    name = f"VMS_Visitor_Report_{loc_part}_{from_s}_to_{to_s}.csv"
    return build_csv(rows), name, len(rows)


def export_visits_xlsx(db: Session, ctx: AuthContext, filters: Dict[str, Any]) -> Tuple[bytes, str, int]:
    _require_export_xlsx(ctx)
    rows = _fetch_export_rows(db, ctx, filters)
    summary = {
        "Report Period": f"{filters['from_dt'].date()} to {filters['to_dt'].date()}",
        "Generated At": datetime.now(timezone.utc).isoformat(),
        "Generated By": ctx.email,
        "Total Visits": len(rows),
    }
    AuditService(db).record(
        action="REPORT_EXPORTED_XLSX",
        entity_type="report",
        entity_id="visits",
        actor_id=ctx.user_id,
        actor_email=ctx.email,
        metadata={"row_count": len(rows)},
    )
    db.commit()
    loc_part = safe_filename_part(str(filters.get("location_label") or "all"))
    from_s = filters["from_dt"].strftime("%Y-%m-%d")
    to_s = filters["to_dt"].strftime("%Y-%m-%d")
    name = f"VMS_Visitor_Report_{loc_part}_{from_s}_to_{to_s}.xlsx"
    return build_xlsx(rows, summary), name, len(rows)
