"""Phase 10: Management dashboard, reporting, and exports."""

import os
import tempfile
from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.application.analytics_service import get_dashboard, list_dashboard_kpi
from app.application.auth_service import build_auth_context
from app.application.bootstrap import bootstrap_development_data
from app.application.report_service import build_csv, sanitize_csv_cell, REPORT_COLUMNS
from app.domain.enums import VisitStatus
from app.domain.models import AdminUser, Host, Location, Visit, Visitor, VisitorType
from app.infrastructure.database import Base, get_db
from app.main import create_app
from tests.test_phase4 import _create_approved_visit


@pytest.fixture
def test_db():
    fd, path = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    engine = create_engine(f"sqlite:///{path}", connect_args={"check_same_thread": False})
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)
    db = TestingSessionLocal()
    bootstrap_development_data(db)
    yield TestingSessionLocal, engine, db
    db.close()
    os.unlink(path)


@pytest.fixture
def client(test_db):
    TestingSessionLocal, _, bootstrap_session = test_db

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app = create_app()
    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as c:
        yield c, TestingSessionLocal, bootstrap_session
    app.dependency_overrides.clear()


def _ctx(session, email: str):
    user = session.query(AdminUser).filter(AdminUser.email == email).first()
    return build_auth_context(user)


def _range_params():
    now = datetime.now(timezone.utc)
    from_dt = (now - timedelta(days=7)).isoformat()
    to_dt = now.isoformat()
    return {"from": from_dt, "to": to_dt}


def test_global_admin_dashboard_access(client):
    c, _, _ = client
    r = c.get("/api/analytics/dashboard", headers={"X-Dev-User-Email": "sheru.bohra@lazypay.in"})
    assert r.status_code == 200
    body = r.json()
    assert "kpis" in body
    assert "visitor_trend" in body


def test_head_admin_dashboard_access(client):
    c, _, session = client
    head = session.query(AdminUser).filter(AdminUser.email == "headadmin@vms.local").first()
    assert head is not None
    r = c.get("/api/analytics/dashboard", headers={"X-Dev-User-Email": "headadmin@vms.local"})
    assert r.status_code == 200


def test_site_admin_dashboard_access(client):
    c, _, _ = client
    r = c.get("/api/analytics/dashboard", headers={"X-Dev-User-Email": "siteadmin.blr@vms.local"})
    assert r.status_code == 200
    body = r.json()
    assert "kpis" in body
    assert "visitor_trend" in body


def test_security_dashboard_access(client):
    c, _, _ = client
    r = c.get("/api/analytics/dashboard", headers={"X-Dev-User-Email": "security.blr@vms.local"})
    assert r.status_code == 200
    body = r.json()
    assert "kpis" in body


def test_site_admin_dashboard_unauthorized_location_denied(client, test_db):
    c, _, session = client
    mum = session.query(Location).filter(Location.code == "MUM").first()
    r = c.get(
        "/api/analytics/dashboard",
        params={"location_id": mum.id, "period": "today"},
        headers={"X-Dev-User-Email": "siteadmin.blr@vms.local"},
    )
    assert r.status_code == 403


def test_security_dashboard_unauthorized_location_denied(client, test_db):
    c, _, session = client
    mum = session.query(Location).filter(Location.code == "MUM").first()
    r = c.get(
        "/api/analytics/dashboard",
        params={"location_id": mum.id, "period": "today"},
        headers={"X-Dev-User-Email": "security.blr@vms.local"},
    )
    assert r.status_code == 403


def test_global_admin_reports_access(client):
    c, _, _ = client
    params = _range_params()
    r = c.get("/api/reports/visits", params=params, headers={"X-Dev-User-Email": "sheru.bohra@lazypay.in"})
    assert r.status_code == 200
    assert "items" in r.json()
    assert "total" in r.json()


def test_head_admin_reports_access(client):
    c, _, session = client
    assert session.query(AdminUser).filter(AdminUser.email == "headadmin@vms.local").first()
    params = _range_params()
    r = c.get("/api/reports/visits", params=params, headers={"X-Dev-User-Email": "headadmin@vms.local"})
    assert r.status_code == 200


def test_site_admin_reports_forbidden(client):
    c, _, _ = client
    params = _range_params()
    r = c.get("/api/reports/visits", params=params, headers={"X-Dev-User-Email": "siteadmin.blr@vms.local"})
    assert r.status_code == 403


def test_security_reports_forbidden(client):
    c, _, _ = client
    params = _range_params()
    r = c.get("/api/reports/visits", params=params, headers={"X-Dev-User-Email": "security.blr@vms.local"})
    assert r.status_code == 403


def test_site_admin_csv_export_forbidden(client):
    c, _, _ = client
    params = _range_params()
    r = c.get("/api/reports/visits/export.csv", params=params, headers={"X-Dev-User-Email": "siteadmin.blr@vms.local"})
    assert r.status_code == 403


def test_security_xlsx_export_forbidden(client):
    c, _, _ = client
    params = _range_params()
    r = c.get("/api/reports/visits/export.xlsx", params=params, headers={"X-Dev-User-Email": "security.blr@vms.local"})
    assert r.status_code == 403


def test_csv_formula_injection_sanitized():
    assert sanitize_csv_cell("=HYPERLINK('evil')").startswith("'")
    assert sanitize_csv_cell("+SUM(1,1)").startswith("'")
    assert sanitize_csv_cell("@CMD").startswith("'")
    full_row = {key: None for key, _ in REPORT_COLUMNS}
    full_row.update({
        "registration_reference": "R1",
        "visitor_name": "=HYPERLINK('x')",
        "company": "+evil",
    })
    content = build_csv([full_row]).decode("utf-8-sig")
    assert "'=HYPERLINK" in content


def test_kpi_counts_match_source(test_db):
    _, _, session = test_db
    owner = _ctx(session, "sheru.bohra@lazypay.in")
    blr = session.query(Location).filter(Location.code == "BLR").first()

    pending_visit, _ = _create_approved_visit(session, "BLR", "kpi-pend")
    pending_visit.status = VisitStatus.PENDING_APPROVAL.value
    session.commit()

    onsite_visit, _ = _create_approved_visit(session, "BLR", "kpi-onsite")
    onsite_visit.status = VisitStatus.ONSITE.value
    onsite_visit.checked_in_at = datetime.now(timezone.utc)
    session.commit()

    data = get_dashboard(session, owner, period="today", location_id=blr.id)
    assert data["kpis"]["awaiting_approval"] >= 1
    assert data["kpis"]["onsite_now"] >= 1


def test_csv_export_generated(client, test_db):
    c, _, session = client
    _create_approved_visit(session, "BLR", "export1")
    params = _range_params()
    r = c.get(
        "/api/reports/visits/export.csv",
        params=params,
        headers={"X-Dev-User-Email": "sheru.bohra@lazypay.in"},
    )
    assert r.status_code == 200
    assert "text/csv" in r.headers.get("content-type", "")
    assert "Visit Reference" in r.text


def test_xlsx_export_generated(client, test_db):
    c, _, session = client
    _create_approved_visit(session, "BLR", "xlsx1")
    params = _range_params()
    r = c.get(
        "/api/reports/visits/export.xlsx",
        params=params,
        headers={"X-Dev-User-Email": "sheru.bohra@lazypay.in"},
    )
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert "Visitor Report" in wb.sheetnames
    ws = wb["Visitor Report"]
    assert ws.cell(1, 1).value == "Visit Reference"


def test_report_location_filter(client, test_db):
    c, _, session = client
    _create_approved_visit(session, "BLR", "blr-only")
    _create_approved_visit(session, "MUM", "mum-only")
    blr = session.query(Location).filter(Location.code == "BLR").first()
    params = _range_params()
    params["location_id"] = blr.id
    r = c.get(
        "/api/reports/visits",
        params=params,
        headers={"X-Dev-User-Email": "sheru.bohra@lazypay.in"},
    )
    assert r.status_code == 200
    for item in r.json()["items"]:
        assert item["location_name"] == blr.name


def test_malicious_visitor_csv_export(client, test_db):
    c, _, session = client
    loc = session.query(Location).filter(Location.code == "BLR").first()
    vt = session.query(VisitorType).filter(VisitorType.code == "BUSINESS").first()
    host = session.query(Host).filter(Host.is_development_seed.is_(True)).first()
    visitor = Visitor(
        full_name="=HYPERLINK('evil')",
        email="evil@test.com",
        phone="+919990000001",
        company="+SUM(1,1)",
        visitor_type_id=vt.id,
    )
    session.add(visitor)
    session.flush()
    visit = Visit(
        visitor_id=visitor.id,
        location_id=loc.id,
        status=VisitStatus.APPROVED.value,
        registration_reference="VMS-EVIL-1",
        host_id=host.id,
        host_name=host.name,
        purpose="Test",
        expected_duration_minutes=60,
        policy_accepted=True,
        source="self_registration",
    )
    session.add(visit)
    session.commit()
    params = _range_params()
    r = c.get(
        "/api/reports/visits/export.csv",
        params=params,
        headers={"X-Dev-User-Email": "sheru.bohra@lazypay.in"},
    )
    assert r.status_code == 200
    # exported cells must be escaped
    assert "'=HYPERLINK" in r.text or "''=HYPERLINK" in r.text


def test_dashboard_location_scope_changes_totals(test_db):
    _, _, session = test_db
    owner = _ctx(session, "sheru.bohra@lazypay.in")
    blr = session.query(Location).filter(Location.code == "BLR").first()
    mum = session.query(Location).filter(Location.code == "MUM").first()
    _create_approved_visit(session, "BLR", "scope-blr")
    _create_approved_visit(session, "MUM", "scope-mum")
    all_data = get_dashboard(session, owner, period="last_7_days")
    blr_data = get_dashboard(session, owner, period="last_7_days", location_id=blr.id)
    mum_data = get_dashboard(session, owner, period="last_7_days", location_id=mum.id)
    all_types = sum(t["count"] for t in all_data["visitor_types"])
    blr_types = sum(t["count"] for t in blr_data["visitor_types"])
    mum_types = sum(t["count"] for t in mum_data["visitor_types"])
    assert all_types >= blr_types
    assert all_types >= mum_types


def test_dashboard_kpi_drill_down_matches_counts(test_db, client):
    c, _, session = client
    owner = _ctx(session, "sheru.bohra@lazypay.in")
    blr = session.query(Location).filter(Location.code == "BLR").first()

    pending_visit, _ = _create_approved_visit(session, "BLR", "drill-pend")
    pending_visit.status = VisitStatus.PENDING_APPROVAL.value
    session.commit()

    now = datetime.now(timezone.utc)
    checked_in_visit, _ = _create_approved_visit(session, "BLR", "drill-checkin")
    checked_in_visit.status = VisitStatus.ONSITE.value
    checked_in_visit.checked_in_at = now
    session.commit()

    checked_out_visit, _ = _create_approved_visit(session, "BLR", "drill-checkout")
    checked_out_visit.status = VisitStatus.CHECKED_OUT.value
    checked_out_visit.checked_in_at = now - timedelta(hours=2)
    checked_out_visit.checked_out_at = now
    session.commit()

    dashboard = get_dashboard(session, owner, period="today", location_id=blr.id)
    headers = {"X-Dev-User-Email": "sheru.bohra@lazypay.in"}

    for key in (
        "awaiting_approval",
        "checked_in_today",
        "onsite_now",
        "checked_out_today",
    ):
        r = c.get(
            f"/api/analytics/dashboard/kpi/{key}",
            params={"location_id": blr.id},
            headers=headers,
        )
        assert r.status_code == 200
        assert r.json()["total"] == dashboard["kpis"][key]

    items, total = list_dashboard_kpi(session, owner, "awaiting_approval", location_id=blr.id)
    assert total == dashboard["kpis"]["awaiting_approval"]
    assert len(items) == min(total, 50)

